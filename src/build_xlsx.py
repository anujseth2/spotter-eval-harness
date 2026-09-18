#!/usr/bin/env python3
"""Build the reviewable accuracy workbook from results/graded.json.

Three sheets:
  Summary  - headline scores (live formulas), accuracy by expected source
  Results  - one row per scored question: question, answer key, verdict, what Spotter
             resolved (tokens + sage query), the gap, latency
  Gaps     - every Partial/Wrong grouped by bucket, which is the improvement roadmap

Optional 4th sheet "Recovered" appears if results/_recovered.json exists (questions
that first failed on the synchronous endpoint and were recovered via streaming).

Needs openpyxl:  pip install openpyxl
"""
import json, os, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import ts_client

ROOT = os.path.join(HERE, "..")
env = ts_client.load_env()

TITLE = env.get("EVAL_TITLE") or "Spotter Accuracy Assessment"
SUBTITLE = env.get("EVAL_SUBTITLE") or \
    "Natural-language questions run through Spotter in AUTO_MODE and graded semantically against a domain answer key."
EXPECTED_MODEL = (env.get("EXPECTED_MODEL_GUID") or "").strip()

# Optional friendly names for the models auto-mode can pick: {"<guid>": "Sales Model", ...}
LABELS_PATH = os.path.join(ROOT, "data/model_labels.json")
MODELS = json.load(open(LABELS_PATH)) if os.path.exists(LABELS_PATH) else {}

graded = json.load(open(os.path.join(ROOT, "results/graded.json")))
qmeta = {str(x["id"]): x for x in json.load(open(os.path.join(ROOT, "data/questions.json")))}

# One representative row per question: the substituted variant when there is one,
# because that is what measures analytical accuracy rather than the "needs context" UX.
by_id = defaultdict(dict)
for g in graded:
    by_id[str(g["id"])][g.get("variant", "cold")] = g
reps = [v.get("substituted") or v["cold"] for v in by_id.values() if ("substituted" in v or "cold" in v)]
reps.sort(key=lambda r: (float(str(r["id"])) if str(r["id"]).replace(".", "", 1).isdigit() else 0))

ARIAL = lambda **k: Font(name="Arial", **k)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
VFILL = {"Correct": "C6EFCE", "Partial": "FFEB9C", "Wrong": "FFC7CE", "Clarification": "DDEBF7",
         "Infra": "D9D9D9", "JudgeError": "E4DFEC"}
VFONT = {"Correct": "006100", "Partial": "9C5700", "Wrong": "9C0006", "Clarification": "1F4E79",
         "Infra": "595959", "JudgeError": "604A7B"}
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------- Results ----------
ws = wb.active; ws.title = "Results"
headers = ["ID", "Bucket", "Expected Source(s)", "Question", "Primary Eval Metrics (answer key)",
           "Expected Answer (answer key)", "Verdict", "Routing OK", "Model Auto-Selected",
           "Resolved Tokens", "Resolved Query (sage)", "Chart Type", "Answer Summary",
           "Gap / Why not Correct", "Run Variant", "Latency (s)", "Retries / Error"]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(1, c); cell.font = ARIAL(bold=True, color="FFFFFF"); cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = BORDER

sub_entity = env.get("SUBSTITUTE_ENTITY", "").strip()
for r in reps:
    qid = str(r["id"]); meta = qmeta.get(qid, {})
    idlabel = qid.replace(".0", "")
    variant = f"entity-substituted ({sub_entity})" if r.get("variant") == "substituted" else "cold"
    guid = r.get("chosen_worksheet_id")
    model = MODELS.get(guid, guid or "")
    if not r.get("produced_answer"):
        routing = ""                                   # routing is meaningless with no answer
    elif EXPECTED_MODEL:
        routing = "Y" if guid == EXPECTED_MODEL else "N"
    else:
        routing = "Y" if r.get("routing_ok") else "N"  # fall back to the judge's call
    toks = ", ".join(r.get("tml_tokens") or []) if r.get("tml_tokens") else ""
    err = r.get("error") or ""
    retry = f"{r.get('attempts')}x" if (r.get("attempts") or 1) > 1 else ""
    retry_err = (retry + (" | " if retry and err else "") + (err[:120] if err else "")).strip()
    narr = (r.get("narrative") or "").replace("\n", " ").strip()
    ws.append([idlabel, r.get("bucket"), r.get("expected_source"), r.get("question"),
               ", ".join(meta.get("primary_eval_metrics") or []), meta.get("expected_answer", ""),
               r.get("verdict"), routing, model, toks, r.get("sage_query") or "",
               r.get("chart_type") or "", narr, r.get("gap") or "", variant,
               r.get("elapsed_s"), retry_err])

wrapcols = {4, 5, 6, 10, 11, 13, 14, 17}
for ri in range(2, ws.max_row + 1):
    v = ws.cell(ri, 7).value
    for c in range(1, len(headers) + 1):
        cell = ws.cell(ri, c); cell.font = ARIAL(size=10); cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(c in wrapcols),
                                   horizontal=("center" if c in (1, 7, 8, 12, 16) else "left"))
    vc = ws.cell(ri, 7)
    if v in VFILL:
        vc.fill = PatternFill("solid", fgColor=VFILL[v]); vc.font = ARIAL(size=10, bold=True, color=VFONT[v])
    rc = ws.cell(ri, 8)
    if rc.value == "N": rc.font = ARIAL(size=10, bold=True, color="9C0006")

widths = [6, 26, 16, 46, 34, 40, 13, 9, 18, 46, 46, 11, 60, 40, 24, 10, 26]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
ws.row_dimensions[1].height = 42
nrows = ws.max_row
ntotal = len(reps)

# ---------- Summary ----------
s = wb.create_sheet("Summary", 0)
def put(cell, val, bold=False, size=11, color="000000", fill=None, align="left", num=None):
    c = s[cell]; c.value = val; c.font = ARIAL(bold=bold, size=size, color=color)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fill: c.fill = PatternFill("solid", fgColor=fill)
    if num: c.number_format = num
    return c

s.merge_cells("A1:D1"); put("A1", TITLE, bold=True, size=16, color="1F3864")
s.merge_cells("A2:D2"); put("A2", SUBTITLE, size=10, color="595959"); s.row_dimensions[2].height = 42
R = "Results"
put("A4", "Headline", bold=True, size=13, color="1F3864")
rows = [("Correct", f'=COUNTIF({R}!G2:G{nrows},"Correct")'),
        ("Partial", f'=COUNTIF({R}!G2:G{nrows},"Partial")'),
        ("Clarification", f'=COUNTIF({R}!G2:G{nrows},"Clarification")'),
        ("Wrong", f'=COUNTIF({R}!G2:G{nrows},"Wrong")'),
        ("Infra (excluded)", f'=COUNTIF({R}!G2:G{nrows},"Infra")'),
        ("Ungraded (JudgeError)", f'=COUNTIF({R}!G2:G{nrows},"JudgeError")')]
r0 = 5
for i, (lab, f) in enumerate(rows):
    put(f"A{r0+i}", lab, bold=True); put(f"B{r0+i}", f, align="center")
    if lab.split()[0] in VFILL:
        s[f"A{r0+i}"].fill = PatternFill("solid", fgColor=VFILL[lab.split()[0]])
put(f"A{r0+6}", "Scored (excl. infra + ungraded)", bold=True)
put(f"B{r0+6}", f"={ntotal}-B{r0+4}-B{r0+5}", align="center")
put(f"A{r0+7}", "Headline accuracy (Correct)", bold=True); put(f"B{r0+7}", f"=B{r0}/B{r0+6}", num="0.0%", align="center")
put(f"A{r0+8}", "Weighted (Correct + 0.5 Partial)", bold=True); put(f"B{r0+8}", f"=(B{r0}+0.5*B{r0+1})/B{r0+6}", num="0.0%", align="center")
put(f"A{r0+9}", "Usable (Correct+Partial+Clarify)", bold=True); put(f"B{r0+9}", f"=(B{r0}+B{r0+1}+B{r0+2})/B{r0+6}", num="0.0%", align="center")
put(f"A{r0+10}", "Auto-mode routing correct", bold=True)
put(f"B{r0+10}", f'=IF(COUNTIF({R}!H2:H{nrows},"Y")+COUNTIF({R}!H2:H{nrows},"N")=0,"-",'
                f'COUNTIF({R}!H2:H{nrows},"Y")/(COUNTIF({R}!H2:H{nrows},"Y")+COUNTIF({R}!H2:H{nrows},"N")))',
    num="0.0%", align="center")

srcs = sorted({r.get("expected_source") or "" for r in reps})
cs = 17
put(f"A{cs}", "Accuracy by expected source", bold=True, size=13, color="1F3864")
for j, h in enumerate(["Source", "n (scored)", "Correct", "Partial", "Wrong", "Clarify", "Correct %"]):
    cc = f"{get_column_letter(1+j)}{cs+1}"; put(cc, h, bold=True, color="FFFFFF", align="center"); s[cc].fill = HDR_FILL
for k, src in enumerate(srcs):
    rr = cs + 2 + k
    put(f"A{rr}", src)
    put(f"B{rr}", f'=COUNTIFS({R}!C2:C{nrows},A{rr},{R}!G2:G{nrows},"<>Infra")', align="center")
    put(f"C{rr}", f'=COUNTIFS({R}!C2:C{nrows},A{rr},{R}!G2:G{nrows},"Correct")', align="center")
    put(f"D{rr}", f'=COUNTIFS({R}!C2:C{nrows},A{rr},{R}!G2:G{nrows},"Partial")', align="center")
    put(f"E{rr}", f'=COUNTIFS({R}!C2:C{nrows},A{rr},{R}!G2:G{nrows},"Wrong")', align="center")
    put(f"F{rr}", f'=COUNTIFS({R}!C2:C{nrows},A{rr},{R}!G2:G{nrows},"Clarification")', align="center")
    put(f"G{rr}", f'=IF(B{rr}=0,"-",C{rr}/B{rr})', num="0.0%", align="center")
for col, w in zip("ABCDEFG", [34, 12, 10, 10, 9, 9, 12]): s.column_dimensions[col].width = w

# ---------- Gaps (the improvement roadmap, derived from the data) ----------
gs = wb.create_sheet("Gaps")
gs.merge_cells("A1:D1")
c = gs["A1"]
c.value = ("Every Partial and Wrong verdict, grouped by bucket. Buckets with the most rows are where "
           "curation or data scope will buy the most accuracy.")
c.font = ARIAL(size=10, italic=True, color="1F3864")
c.alignment = Alignment(wrap_text=True, vertical="center"); gs.row_dimensions[1].height = 32
gs.append(["Bucket", "ID", "Verdict", "Question", "Gap"])
for cc in range(1, 6):
    cell = gs.cell(2, cc); cell.font = ARIAL(bold=True, color="FFFFFF"); cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

shortfall = [r for r in reps if r.get("verdict") in ("Partial", "Wrong")]
bucket_order = sorted({r.get("bucket") or "" for r in shortfall},
                      key=lambda b: -sum(1 for r in shortfall if (r.get("bucket") or "") == b))
for b in bucket_order:
    for r in [x for x in shortfall if (x.get("bucket") or "") == b]:
        gs.append([b, str(r["id"]).replace(".0", ""), r.get("verdict"), r.get("question"), r.get("gap") or ""])
for ri in range(3, gs.max_row + 1):
    for cc in range(1, 6):
        gs.cell(ri, cc).font = ARIAL(size=10)
        gs.cell(ri, cc).alignment = Alignment(vertical="top", wrap_text=(cc in (4, 5)))
    vv = gs.cell(ri, 3).value
    if vv in VFILL:
        gs.cell(ri, 3).fill = PatternFill("solid", fgColor=VFILL[vv])
        gs.cell(ri, 3).font = ARIAL(size=10, bold=True, color=VFONT[vv])
for col, w in zip("ABCDE", [28, 6, 12, 54, 60]): gs.column_dimensions[col].width = w
gs.freeze_panes = "A3"

# ---------- Recovered (only when a recovery re-run happened) ----------
rec_path = os.path.join(ROOT, "results/_recovered.json")
if os.path.exists(rec_path):
    raw = json.load(open(rec_path))
    rec = raw if isinstance(raw, dict) else {str(x["id"]): x for x in raw}
    rl = wb.create_sheet("Recovered")
    rl.merge_cells("A1:G1")
    note = ("These queries first failed with a gateway timeout on the synchronous REST endpoint under "
            "concurrency. They are not a Spotter comprehension limit: re-running them serially and via the "
            "streaming endpoint the Spotter UI uses recovered them. Their verdicts are folded into the main scores.")
    c = rl["A1"]; c.value = note; c.font = ARIAL(size=10, italic=True, color="1F3864")
    c.alignment = Alignment(wrap_text=True, vertical="center"); rl.row_dimensions[1].height = 54
    rl.append(["ID", "Bucket", "Question", "Transport that recovered it", "Latency (s)", "Verdict", "Resolved Query (sage)"])
    for cc in range(1, 8):
        cell = rl.cell(2, cc); cell.font = ARIAL(bold=True, color="FFFFFF"); cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    byid_reps = {str(r["id"]): r for r in reps}
    for k in sorted(rec, key=lambda t: float(t) if str(t).replace(".", "", 1).isdigit() else 0):
        x = rec[k]; rep = byid_reps.get(k, {})
        transport = "streaming (UI transport)" if x.get("transport") == "stream" else "serial (synchronous)"
        rl.append([str(k).replace(".0", ""), x.get("bucket"), x.get("question"), transport,
                   x.get("elapsed_s"), rep.get("verdict"), x.get("sage_query")])
    for ri in range(3, rl.max_row + 1):
        for cc in range(1, 8):
            rl.cell(ri, cc).font = ARIAL(size=10)
            rl.cell(ri, cc).alignment = Alignment(vertical="top", wrap_text=(cc in (3, 7)))
        vv = rl.cell(ri, 6).value
        if vv in VFILL:
            rl.cell(ri, 6).fill = PatternFill("solid", fgColor=VFILL[vv])
            rl.cell(ri, 6).font = ARIAL(size=10, bold=True, color=VFONT[vv])
    for col, w in zip("ABCDEFG", [6, 24, 46, 26, 10, 13, 46]): rl.column_dimensions[col].width = w
    rl.freeze_panes = "A3"

out = os.path.join(ROOT, "results", (env.get("EVAL_SLUG") or "spotter") + "_accuracy_results.xlsx")
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print("saved", out, "| scored rows:", ntotal, "| gaps:", len(shortfall))
